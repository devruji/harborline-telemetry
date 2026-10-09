"""Logistics spreadsheet parser."""

from pathlib import Path

from openpyxl import load_workbook


def records(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    parsed = []
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue
        parsed.append(dict(zip(header, row)))
    workbook.close()
    return parsed
